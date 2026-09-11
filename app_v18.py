"""
Print3D Analyzer + Gestión de Proyectos
Backend único en FastAPI. Persistencia en SQLite (/opt/print3d/data.db).
"""
from fastapi import FastAPI, UploadFile, File, Query, Request, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from starlette.middleware.base import BaseHTTPMiddleware
import trimesh, tempfile, os, json, sqlite3, datetime, urllib.parse, io, hashlib, secrets
import cadquery as cq
import requests

app = FastAPI(title="Print3D Analyzer + Proyectos")

class PrivateNetworkMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        if request.method == "OPTIONS":
            from starlette.responses import Response
            r = Response()
            r.headers["Access-Control-Allow-Private-Network"] = "true"
            r.headers["Access-Control-Allow-Origin"] = "*"
            r.headers["Access-Control-Allow-Headers"] = "*"
            return r
        response = await call_next(request)
        response.headers["Access-Control-Allow-Private-Network"] = "true"
        return response

app.add_middleware(PrivateNetworkMiddleware)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

CONFIG_FILE = "/opt/print3d/config.json"
DB_FILE = "/opt/print3d/data.db"

DEFAULT_CONFIG = {
    "materials": [
        {"id":"pla",   "name":"PLA",   "price":20, "density":1.24, "type":"fdm"},
        {"id":"pla+",  "name":"PLA+",  "price":24, "density":1.24, "type":"fdm"},
        {"id":"petg",  "name":"PETG",  "price":25, "density":1.27, "type":"fdm"},
        {"id":"abs",   "name":"ABS",   "price":22, "density":1.04, "type":"fdm"},
        {"id":"asa",   "name":"ASA",   "price":28, "density":1.07, "type":"fdm"},
        {"id":"tpu",   "name":"TPU",   "price":40, "density":1.21, "type":"fdm"},
        {"id":"nylon", "name":"Nylon", "price":45, "density":1.13, "type":"fdm"},
        {"id":"resin", "name":"Resina","price":35, "density":1.10, "type":"resin"}
    ],
    "printers": [
        {"id":"bambu",  "name":"Bambu H2D",        "speed":250, "layer":0.2,  "watts":450},
        {"id":"ender",  "name":"Ender 3",           "speed":60,  "layer":0.2,  "watts":150},
        {"id":"ratrig", "name":"RatRig V-Core 500", "speed":80,  "layer":0.2,  "watts":600},
        {"id":"jupiter","name":"Jupiter SE",         "speed":0,   "layer":0.05, "watts":50}
    ],
    "defaults": {"labor_cost": 15, "margin": 30, "kwh_price": 0.18, "iva": 21}
}

TIPOS_TRABAJO = ["IR", "D", "DO", "I3D", "MD", "FA"]
ESTADOS_PROYECTO = ["presupuestado", "aceptado", "en_curso", "entregado", "facturado", "cobrado", "rechazado"]
MARCAS = ["myrox_lab", "myrox_print", "myrox_works", "myrox_automation", "mw3d"]

# ──────────────────────────────────────────────────────────────────
# DB SETUP
# ──────────────────────────────────────────────────────────────────

def get_db():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()

    # ---- Presupuestadora: trabajos / stock / consumibles / amortización ----
    c.execute("""CREATE TABLE IF NOT EXISTS trabajos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        fecha TEXT DEFAULT (datetime('now')),
        nombre TEXT, material TEXT, impresora TEXT,
        volumen_cm3 REAL, tiempo_estimado_h REAL, tiempo_real_h REAL,
        precio_sin_iva REAL, precio_con_iva REAL, precio_cobrado REAL,
        ganancia REAL, fallo TEXT, notas TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS amortizacion (
        printer_id TEXT PRIMARY KEY,
        coste_compra REAL, vida_util_h REAL, mantenimiento_mes REAL,
        horas_registradas REAL DEFAULT 0
    )""")

    # Costes fijos mensuales del taller (fila única) para calcular un coste estructural por hora
    c.execute("""CREATE TABLE IF NOT EXISTS costes_fijos (
        id INTEGER PRIMARY KEY CHECK (id = 1),
        alquiler_nave REAL DEFAULT 0,
        electricidad REAL DEFAULT 0,
        cuota_autonomos REAL DEFAULT 0,
        reparaciones REAL DEFAULT 0,
        inversiones REAL DEFAULT 0,
        horas_taller_mes REAL DEFAULT 0
    )""")

    # ---- Login por persona ----
    # 'rol' existe ya de cara al futuro: hoy los 3 perfiles son 'admin' (mismo acceso),
    # pero deja preparado el hueco para roles diferenciados (p.ej. 'agente' para un
    # futuro agente automatizado, o 'lectura' para un perfil de solo consulta) sin
    # tener que tocar el esquema otra vez.
    c.execute("""CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT NOT NULL,
        login TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        salt TEXT NOT NULL,
        rol TEXT NOT NULL DEFAULT 'admin',
        activo INTEGER NOT NULL DEFAULT 1,
        debe_cambiar_password INTEGER NOT NULL DEFAULT 1,
        creado_en TEXT DEFAULT (datetime('now'))
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS sesiones (
        token TEXT PRIMARY KEY,
        usuario_id INTEGER NOT NULL REFERENCES usuarios(id),
        creado_en TEXT DEFAULT (datetime('now')),
        expira_en TEXT NOT NULL
    )""")

    # Fotos adjuntas y comentarios en las tarjetas del tablero (kanban)
    c.execute("""CREATE TABLE IF NOT EXISTS proyecto_fotos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        proyecto_id INTEGER NOT NULL REFERENCES proyectos(id),
        filename TEXT NOT NULL,
        usuario TEXT,
        fecha TEXT DEFAULT (datetime('now'))
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS proyecto_comentarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        proyecto_id INTEGER NOT NULL REFERENCES proyectos(id),
        usuario TEXT,
        texto TEXT NOT NULL,
        fecha TEXT DEFAULT (datetime('now'))
    )""")

    # Albaranes de entrega — pueden nacer de un presupuesto aceptado (heredan sus líneas) o crearse sueltos
    c.execute("""CREATE TABLE IF NOT EXISTS albaranes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        proyecto_id INTEGER NOT NULL REFERENCES proyectos(id),
        presupuesto_id INTEGER REFERENCES presupuestos(id),
        numero TEXT UNIQUE NOT NULL,
        fecha TEXT DEFAULT (datetime('now')),
        estado TEXT DEFAULT 'emitido',
        notas TEXT,
        creado_por TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS albaran_lineas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        albaran_id INTEGER NOT NULL REFERENCES albaranes(id),
        descripcion TEXT,
        cantidad REAL
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS consumibles (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        fecha TEXT, printer_id TEXT, tipo TEXT, descripcion TEXT, coste REAL
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS stock_materiales (
        material_id TEXT PRIMARY KEY,
        cantidad REAL DEFAULT 0,
        stock_minimo REAL DEFAULT 0,
        precio_medio REAL DEFAULT 0
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS stock_movimientos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        fecha TEXT DEFAULT (datetime('now')),
        material_id TEXT, tipo TEXT, cantidad REAL, coste_total REAL, notas TEXT
    )""")

    # ---- Módulo de Proyectos / CRM ----
    c.execute("""CREATE TABLE IF NOT EXISTS clientes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT NOT NULL, empresa TEXT,
        contacto_email TEXT, contacto_telefono TEXT,
        fecha_alta TEXT DEFAULT (datetime('now'))
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS contactos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_id INTEGER NOT NULL REFERENCES clientes(id),
        nombre TEXT NOT NULL, email TEXT, telefono TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS contadores_codigo (
        tipo TEXT NOT NULL, anio INTEGER NOT NULL,
        ultimo_numero INTEGER DEFAULT 0,
        PRIMARY KEY (tipo, anio)
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS proyectos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo TEXT UNIQUE NOT NULL,
        nombre TEXT,
        cliente_id INTEGER REFERENCES clientes(id),
        contacto_id INTEGER REFERENCES contactos(id),
        marca TEXT, tipo_trabajo TEXT,
        estado TEXT DEFAULT 'presupuestado',
        fecha_alta TEXT DEFAULT (datetime('now')),
        fecha_entrega_compromiso TEXT, fecha_entrega_real TEXT,
        fecha_facturacion TEXT, fecha_cobro TEXT,
        ruta_nextcloud TEXT, imagen_portada_ruta TEXT, notas TEXT,
        archivado INTEGER DEFAULT 0
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS presupuestos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        proyecto_id INTEGER NOT NULL REFERENCES proyectos(id),
        version INTEGER, fecha TEXT DEFAULT (datetime('now')),
        estado TEXT DEFAULT 'borrador', importe REAL DEFAULT 0,
        referencia_archivo TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS presupuesto_lineas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        presupuesto_id INTEGER NOT NULL REFERENCES presupuestos(id),
        tipo TEXT,
        descripcion TEXT,
        cantidad REAL,
        precio_unitario REAL,
        importe REAL
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS tarifas_horas (
        tipo_hora TEXT PRIMARY KEY,
        precio_hora REAL
    )""")
    for tipo, precio in [("ir", 35), ("diseno", 30), ("diseno_organico", 35),
                          ("impresion", 15), ("reunion", 25), ("montaje", 25), ("km", 0.30)]:
        c.execute("INSERT OR IGNORE INTO tarifas_horas (tipo_hora, precio_hora) VALUES (?,?)", (tipo, precio))

    c.execute("""CREATE TABLE IF NOT EXISTS facturas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        proyecto_id INTEGER NOT NULL REFERENCES proyectos(id),
        numero TEXT, importe REAL, fecha TEXT,
        estado_cobro TEXT, referencia_archivo TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS horas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        proyecto_id INTEGER NOT NULL REFERENCES proyectos(id),
        fecha TEXT, tipo_trabajo_hora TEXT, horas REAL, descripcion TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS gastos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        proyecto_id INTEGER NOT NULL REFERENCES proyectos(id),
        fecha TEXT, descripcion TEXT, importe REAL
    )""")

    # Migración: añadir columna archivado si no existe
    try:
        c.execute("ALTER TABLE proyectos ADD COLUMN archivado INTEGER DEFAULT 0")
        conn.commit()
    except:
        pass

    # Migración: trazabilidad de "quién hizo cada acción" (3 perfiles, mismo nivel de acceso)
    for tabla, columna in [
        ("proyectos", "creado_por"),
        ("presupuestos", "creado_por"),
        ("trabajos", "usuario"),
        ("consumibles", "usuario"),
        ("stock_movimientos", "usuario"),
        ("horas", "usuario"),
        ("gastos", "usuario"),
    ]:
        try:
            c.execute(f"ALTER TABLE {tabla} ADD COLUMN {columna} TEXT")
            conn.commit()
        except:
            pass

    # Migración: prioridad del proyecto (franja de color en la tarjeta del tablero)
    try:
        c.execute("ALTER TABLE proyectos ADD COLUMN prioridad TEXT DEFAULT 'normal'")
        conn.commit()
    except:
        pass

    conn.commit()
    conn.close()
init_db()

def _hash_password(password: str, salt: str = None):
    salt = salt or secrets.token_hex(16)
    h = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 200_000)
    return h.hex(), salt

def _verificar_password(password: str, password_hash: str, salt: str) -> bool:
    h, _ = _hash_password(password, salt)
    return secrets.compare_digest(h, password_hash)

def _seed_usuarios():
    """Crea las 3 cuentas iniciales la primera vez que arranca, con contraseña temporal.
    Cada una debe cambiarla en su primer login (debe_cambiar_password=1)."""
    conn = get_db()
    n = conn.execute("SELECT COUNT(*) as n FROM usuarios").fetchone()["n"]
    if n == 0:
        iniciales = [
            ("Christian", "christian", "christian2026"),
            ("Santi",     "santi",     "santi2026"),
            ("Belén",     "belen",     "belen2026"),
        ]
        for nombre, login, temp_pass in iniciales:
            h, salt = _hash_password(temp_pass)
            conn.execute("""INSERT INTO usuarios (nombre, login, password_hash, salt, rol, activo, debe_cambiar_password)
                VALUES (?,?,?,?,'admin',1,1)""", (nombre, login, h, salt))
        conn.commit()
    conn.close()
_seed_usuarios()

SESSION_DURATION_HORAS = 24 * 14  # dos semanas

def _crear_sesion(usuario_id: int) -> str:
    token = secrets.token_urlsafe(32)
    expira = (datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=SESSION_DURATION_HORAS)).isoformat()
    conn = get_db()
    conn.execute("INSERT INTO sesiones (token, usuario_id, expira_en) VALUES (?,?,?)", (token, usuario_id, expira))
    conn.commit()
    conn.close()
    return token

def _usuario_de_sesion(request: Request):
    """Resuelve el usuario autenticado a partir del token de sesión (cabecera
    'Authorization: Bearer <token>'). Devuelve None si no hay sesión válida.
    A diferencia del esquema anterior, el nombre NO viene de una cabecera que el
    cliente pueda inventarse — sale de la sesión creada en /auth/login."""
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None
    token = auth[len("Bearer "):].strip()
    if not token:
        return None
    conn = get_db()
    row = conn.execute("""
        SELECT u.id, u.nombre, u.login, u.rol, u.activo, u.debe_cambiar_password, s.expira_en
        FROM sesiones s JOIN usuarios u ON u.id = s.usuario_id
        WHERE s.token = ?
    """, (token,)).fetchone()
    conn.close()
    if not row or not row["activo"]:
        return None
    if row["expira_en"] < datetime.datetime.now(datetime.timezone.utc).isoformat():
        return None
    return dict(row)

def _usuario(request: Request) -> str:
    """Nombre de quien hizo la acción (para trazabilidad), a partir de la sesión autenticada."""
    u = _usuario_de_sesion(request)
    return u["nombre"] if u else None

def _requerir_usuario(request: Request) -> dict:
    """Como _usuario_de_sesion, pero exige que haya sesión válida (401 si no)."""
    u = _usuario_de_sesion(request)
    if not u:
        raise HTTPException(401, "Sesión no válida o caducada")
    return u

# ──────────────────────────────────────────────────────────────────
# LOGIN POR PERSONA
# ──────────────────────────────────────────────────────────────────

@app.post("/auth/login")
async def login(request: Request):
    d = await request.json()
    login_ = (d.get("login") or "").strip().lower()
    password = d.get("password") or ""
    conn = get_db()
    row = conn.execute("SELECT * FROM usuarios WHERE login=?", (login_,)).fetchone()
    conn.close()
    if not row or not row["activo"] or not _verificar_password(password, row["password_hash"], row["salt"]):
        raise HTTPException(401, "Usuario o contraseña incorrectos")
    token = _crear_sesion(row["id"])
    return {
        "token": token,
        "nombre": row["nombre"],
        "rol": row["rol"],
        "debe_cambiar_password": bool(row["debe_cambiar_password"]),
    }

@app.get("/auth/me")
def auth_me(request: Request):
    u = _requerir_usuario(request)
    return {"nombre": u["nombre"], "rol": u["rol"], "debe_cambiar_password": bool(u["debe_cambiar_password"])}

@app.post("/auth/logout")
def logout(request: Request):
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        token = auth[len("Bearer "):].strip()
        conn = get_db()
        conn.execute("DELETE FROM sesiones WHERE token=?", (token,))
        conn.commit()
        conn.close()
    return {"status": "ok"}

@app.post("/auth/cambiar-password")
async def cambiar_password(request: Request):
    u = _requerir_usuario(request)
    d = await request.json()
    password_actual = d.get("password_actual") or ""
    password_nueva = d.get("password_nueva") or ""
    if len(password_nueva) < 6:
        raise HTTPException(400, "La contraseña nueva debe tener al menos 6 caracteres")
    conn = get_db()
    row = conn.execute("SELECT * FROM usuarios WHERE id=?", (u["id"],)).fetchone()
    if not _verificar_password(password_actual, row["password_hash"], row["salt"]):
        conn.close()
        raise HTTPException(401, "La contraseña actual no es correcta")
    h, salt = _hash_password(password_nueva)
    conn.execute("UPDATE usuarios SET password_hash=?, salt=?, debe_cambiar_password=0 WHERE id=?", (h, salt, u["id"]))
    conn.commit()
    conn.close()
    return {"status": "ok"}

# ---- Gestión de usuarios (pensado para cuando quieras segregar accesos o dar de
# alta una cuenta para un agente automatizado: crea la cuenta aquí con el rol que
# quieras, p.ej. 'agente', y usa /auth/login con esas credenciales igual que un
# usuario humano — no hace falta tocar el resto del backend) ----

@app.get("/usuarios")
def listar_usuarios(request: Request):
    _requerir_usuario(request)
    conn = get_db()
    rows = conn.execute("SELECT id, nombre, login, rol, activo, creado_en FROM usuarios ORDER BY id").fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/usuarios")
async def crear_usuario(request: Request):
    _requerir_usuario(request)
    d = await request.json()
    nombre = (d.get("nombre") or "").strip()
    login_ = (d.get("login") or "").strip().lower()
    password = d.get("password") or ""
    rol = d.get("rol") or "admin"
    if not nombre or not login_ or len(password) < 6:
        raise HTTPException(400, "Faltan datos o la contraseña es demasiado corta")
    h, salt = _hash_password(password)
    conn = get_db()
    try:
        cur = conn.execute("""INSERT INTO usuarios (nombre, login, password_hash, salt, rol, activo, debe_cambiar_password)
            VALUES (?,?,?,?,?,1,1)""", (nombre, login_, h, salt, rol))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        raise HTTPException(400, f"Ya existe un usuario con el login '{login_}'")
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.patch("/usuarios/{usuario_id}")
async def editar_usuario(usuario_id: int, request: Request):
    _requerir_usuario(request)
    d = await request.json()
    conn = get_db()
    if "activo" in d:
        conn.execute("UPDATE usuarios SET activo=? WHERE id=?", (1 if d["activo"] else 0, usuario_id))
    if "rol" in d:
        conn.execute("UPDATE usuarios SET rol=? WHERE id=?", (d["rol"], usuario_id))
    if "reset_password" in d and d["reset_password"]:
        h, salt = _hash_password(d["reset_password"])
        conn.execute("UPDATE usuarios SET password_hash=?, salt=?, debe_cambiar_password=1 WHERE id=?", (h, salt, usuario_id))
    conn.commit()
    conn.close()
    return {"status": "saved"}

def get_or_create_config():
    if not os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, "w") as f:
            json.dump(DEFAULT_CONFIG, f, indent=2)
    with open(CONFIG_FILE, "r") as f:
        return json.load(f)

def load_step(path: str):
    shape = cq.importers.importStep(path)
    bb = shape.val().BoundingBox()
    volume_mm3 = shape.val().Volume()
    bbox = [round(bb.xmax - bb.xmin, 2), round(bb.ymax - bb.ymin, 2), round(bb.zmax - bb.zmin, 2)]
    return abs(volume_mm3), bbox

def row_to_dict(row):
    return dict(row) if row else None

# ──────────────────────────────────────────────────────────────────
# HEALTH / CONFIG / ANALYZE (igual que antes)
# ──────────────────────────────────────────────────────────────────

@app.get("/health")
def health():
    return {"status": "ok"}

@app.get("/config")
def get_config():
    return get_or_create_config()

@app.post("/config")
async def save_config(request: Request):
    data = await request.json()
    with open(CONFIG_FILE, "w") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    return {"status": "saved"}

@app.post("/analyze")
async def analyze(
    file: UploadFile = File(...),
    material: str = Query("pla"),
    infill: float = Query(0.20),
    fail_rate: float = Query(0.05),
    margin: float = Query(0.30),
    machine_cost_h: float = Query(0.50),
    labor_cost_h: float = Query(5.0),
    qty: int = Query(1),
):
    suffix = os.path.splitext(file.filename)[1].lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name
    try:
        if suffix in (".step", ".stp"):
            volume_mm3, bbox = load_step(tmp_path)
            is_watertight, faces = True, None
        else:
            mesh = trimesh.load(tmp_path)
            if isinstance(mesh, trimesh.Scene):
                mesh = trimesh.util.concatenate(list(mesh.geometry.values()))
            volume_mm3 = abs(mesh.volume)
            bbox = [round(x, 1) for x in mesh.bounding_box.extents]
            is_watertight = bool(mesh.is_watertight)
            faces = len(mesh.faces)
    finally:
        os.unlink(tmp_path)

    volume_cm3 = volume_mm3 / 1000
    config = get_or_create_config()
    mat = next((m for m in config["materials"] if m["id"] == material), config["materials"][0])

    shell_fraction = 0.25
    eff_vol = volume_cm3 * (shell_fraction + (1 - shell_fraction) * infill)
    weight_g = eff_vol * mat["density"]
    mat_cost = (weight_g / 1000) * mat["price"]
    print_time_h = (volume_cm3 * 0.12) / (infill + 0.3)
    machine_cost = print_time_h * machine_cost_h
    labor_cost = (print_time_h * 0.2 + 0.5) * labor_cost_h
    fail_coeff = 1 / (1 - fail_rate) if fail_rate < 1 else 2.0
    subtotal = (mat_cost + machine_cost + labor_cost) * fail_coeff
    price_unit = subtotal * (1 + margin)
    price_total = price_unit * qty

    return {
        "file": file.filename, "material": mat["name"],
        "geometry": {"volume_cm3": round(volume_cm3, 3), "weight_g": round(weight_g, 2),
                     "bounding_box_mm": bbox, "is_watertight": is_watertight, "faces": faces, "format": suffix},
        "print": {"infill_pct": round(infill * 100), "estimated_time_h": round(print_time_h, 2)},
        "costs": {"material": round(mat_cost, 3), "machine": round(machine_cost, 3), "labor": round(labor_cost, 3),
                  "fail_coeff": round(fail_coeff, 4), "margin_pct": round(margin * 100)},
        "price_unit": round(price_unit, 2), "price_total": round(price_total, 2), "qty": qty,
    }

# ──────────────────────────────────────────────────────────────────
# TRABAJOS / HISTORIAL (persistencia real)
# ──────────────────────────────────────────────────────────────────

@app.post("/trabajos")
async def crear_trabajo(request: Request):
    d = await request.json()
    conn = get_db()
    cur = conn.execute("""INSERT INTO trabajos
        (nombre, material, impresora, volumen_cm3, tiempo_estimado_h, tiempo_real_h,
         precio_sin_iva, precio_con_iva, precio_cobrado, ganancia, fallo, notas, usuario)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (d.get("nombre"), d.get("material"), d.get("impresora"), d.get("volumen_cm3"),
         d.get("tiempo_estimado_h"), d.get("tiempo_real_h"), d.get("precio_sin_iva"),
         d.get("precio_con_iva"), d.get("precio_cobrado"), d.get("ganancia"),
         d.get("fallo"), d.get("notas"), _usuario(request)))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.get("/trabajos")
def listar_trabajos():
    conn = get_db()
    rows = conn.execute("SELECT * FROM trabajos ORDER BY fecha DESC").fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ──────────────────────────────────────────────────────────────────
# AMORTIZACIÓN
# ──────────────────────────────────────────────────────────────────

@app.post("/amortizacion")
async def guardar_amortizacion(request: Request):
    d = await request.json()
    conn = get_db()
    conn.execute("""INSERT INTO amortizacion (printer_id, coste_compra, vida_util_h, mantenimiento_mes)
        VALUES (?,?,?,?)
        ON CONFLICT(printer_id) DO UPDATE SET
        coste_compra=excluded.coste_compra, vida_util_h=excluded.vida_util_h,
        mantenimiento_mes=excluded.mantenimiento_mes""",
        (d["printer_id"], d.get("coste_compra"), d.get("vida_util_h"), d.get("mantenimiento_mes")))
    conn.commit()
    conn.close()
    return {"status": "saved"}

@app.get("/amortizacion")
def listar_amortizacion():
    conn = get_db()
    rows = conn.execute("SELECT * FROM amortizacion").fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ──────────────────────────────────────────────────────────────────
# COSTES FIJOS (coste estructural por hora)
# ──────────────────────────────────────────────────────────────────

@app.get("/costes-fijos")
def obtener_costes_fijos():
    conn = get_db()
    row = conn.execute("SELECT * FROM costes_fijos WHERE id=1").fetchone()
    conn.close()
    if not row:
        data = {"alquiler_nave": 0, "electricidad": 0, "cuota_autonomos": 0,
                "reparaciones": 0, "inversiones": 0, "horas_taller_mes": 0}
    else:
        data = dict(row)
    total_mes = (data.get("alquiler_nave", 0) + data.get("electricidad", 0) +
                 data.get("cuota_autonomos", 0) + data.get("reparaciones", 0) +
                 data.get("inversiones", 0))
    horas = data.get("horas_taller_mes") or 0
    data["total_mes"] = round(total_mes, 2)
    data["coste_hora"] = round(total_mes / horas, 2) if horas > 0 else 0
    return data

@app.post("/costes-fijos")
async def guardar_costes_fijos(request: Request):
    d = await request.json()
    conn = get_db()
    conn.execute("""INSERT INTO costes_fijos
        (id, alquiler_nave, electricidad, cuota_autonomos, reparaciones, inversiones, horas_taller_mes)
        VALUES (1,?,?,?,?,?,?)
        ON CONFLICT(id) DO UPDATE SET
        alquiler_nave=excluded.alquiler_nave, electricidad=excluded.electricidad,
        cuota_autonomos=excluded.cuota_autonomos, reparaciones=excluded.reparaciones,
        inversiones=excluded.inversiones, horas_taller_mes=excluded.horas_taller_mes""",
        (d.get("alquiler_nave", 0), d.get("electricidad", 0), d.get("cuota_autonomos", 0),
         d.get("reparaciones", 0), d.get("inversiones", 0), d.get("horas_taller_mes", 0)))
    conn.commit()
    conn.close()
    return {"status": "saved"}

# ──────────────────────────────────────────────────────────────────
# CONSUMIBLES
# ──────────────────────────────────────────────────────────────────

@app.post("/consumibles")
async def crear_consumible(request: Request):
    d = await request.json()
    conn = get_db()
    conn.execute("""INSERT INTO consumibles (fecha, printer_id, tipo, descripcion, coste, usuario)
        VALUES (?,?,?,?,?,?)""",
        (d.get("fecha"), d.get("printer_id"), d.get("tipo"), d.get("descripcion"), d.get("coste"), _usuario(request)))
    conn.commit()
    conn.close()
    return {"status": "saved"}

@app.get("/consumibles")
def listar_consumibles():
    conn = get_db()
    rows = conn.execute("SELECT * FROM consumibles ORDER BY fecha DESC").fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ──────────────────────────────────────────────────────────────────
# STOCK
# ──────────────────────────────────────────────────────────────────

@app.post("/stock/entrada")
async def stock_entrada(request: Request):
    d = await request.json()
    mat_id, qty, coste = d["material_id"], float(d["cantidad"]), float(d.get("coste_total") or 0)
    conn = get_db()
    row = conn.execute("SELECT * FROM stock_materiales WHERE material_id=?", (mat_id,)).fetchone()
    if row:
        nueva_cant = row["cantidad"] + qty
        valor_actual = row["cantidad"] * row["precio_medio"]
        nuevo_precio_medio = (valor_actual + coste) / nueva_cant if nueva_cant else 0
        conn.execute("UPDATE stock_materiales SET cantidad=?, precio_medio=? WHERE material_id=?",
                     (nueva_cant, nuevo_precio_medio, mat_id))
    else:
        precio_medio = coste / qty if qty else 0
        conn.execute("INSERT INTO stock_materiales (material_id, cantidad, precio_medio) VALUES (?,?,?)",
                     (mat_id, qty, precio_medio))
    conn.execute("INSERT INTO stock_movimientos (material_id, tipo, cantidad, coste_total, notas, usuario) VALUES (?,?,?,?,?,?)",
                 (mat_id, "entrada", qty, coste, d.get("notas"), _usuario(request)))
    conn.commit()
    conn.close()
    return {"status": "saved"}

@app.post("/stock/salida")
async def stock_salida(request: Request):
    d = await request.json()
    mat_id, qty = d["material_id"], float(d["cantidad"])
    conn = get_db()
    row = conn.execute("SELECT * FROM stock_materiales WHERE material_id=?", (mat_id,)).fetchone()
    nueva_cant = (row["cantidad"] - qty) if row else -qty
    if row:
        conn.execute("UPDATE stock_materiales SET cantidad=? WHERE material_id=?", (nueva_cant, mat_id))
    else:
        conn.execute("INSERT INTO stock_materiales (material_id, cantidad) VALUES (?,?)", (mat_id, nueva_cant))
    conn.execute("INSERT INTO stock_movimientos (material_id, tipo, cantidad, notas, usuario) VALUES (?,?,?,?,?)",
                 (mat_id, "salida", qty, d.get("notas"), _usuario(request)))
    conn.commit()
    conn.close()
    return {"status": "saved"}

@app.get("/stock")
def listar_stock():
    conn = get_db()
    rows = conn.execute("SELECT * FROM stock_materiales").fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ──────────────────────────────────────────────────────────────────
# CLIENTES / CONTACTOS
# ──────────────────────────────────────────────────────────────────

@app.post("/clientes")
async def crear_cliente(request: Request):
    d = await request.json()
    conn = get_db()
    cur = conn.execute("INSERT INTO clientes (nombre, empresa, contacto_email, contacto_telefono) VALUES (?,?,?,?)",
                        (d["nombre"], d.get("empresa"), d.get("contacto_email"), d.get("contacto_telefono")))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.get("/clientes")
def listar_clientes():
    conn = get_db()
    rows = conn.execute("SELECT * FROM clientes ORDER BY nombre").fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/clientes/{cliente_id}")
def detalle_cliente(cliente_id: int):
    conn = get_db()
    row = conn.execute("SELECT * FROM clientes WHERE id=?", (cliente_id,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Cliente no encontrado")
    return dict(row)

@app.post("/clientes/{cliente_id}/contactos")
async def crear_contacto(cliente_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    cur = conn.execute("INSERT INTO contactos (cliente_id, nombre, email, telefono) VALUES (?,?,?,?)",
                        (cliente_id, d["nombre"], d.get("email"), d.get("telefono")))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.get("/clientes/{cliente_id}/contactos")
def listar_contactos(cliente_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM contactos WHERE cliente_id=? ORDER BY nombre", (cliente_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ──────────────────────────────────────────────────────────────────
# PROYECTOS
# ──────────────────────────────────────────────────────────────────

def generar_codigo(conn, tipo: str) -> str:
    if tipo not in TIPOS_TRABAJO:
        raise HTTPException(400, f"Tipo de trabajo inválido: {tipo}")
    anio = datetime.date.today().year
    conn.execute("""INSERT INTO contadores_codigo (tipo, anio, ultimo_numero) VALUES (?,?,1)
        ON CONFLICT(tipo, anio) DO UPDATE SET ultimo_numero = ultimo_numero + 1""", (tipo, anio))
    row = conn.execute("SELECT ultimo_numero FROM contadores_codigo WHERE tipo=? AND anio=?", (tipo, anio)).fetchone()
    numero = row["ultimo_numero"]
    return f"{tipo}-{anio}-{numero:03d}"

def generar_numero_albaran(conn) -> str:
    """Mismo mecanismo de numeración que generar_codigo, pero con su propia serie 'ALB'
    (independiente de TIPOS_TRABAJO), formato ALB-2026-001 — igual que ya usáis a mano."""
    anio = datetime.date.today().year
    conn.execute("""INSERT INTO contadores_codigo (tipo, anio, ultimo_numero) VALUES ('ALB',?,1)
        ON CONFLICT(tipo, anio) DO UPDATE SET ultimo_numero = ultimo_numero + 1""", (anio,))
    row = conn.execute("SELECT ultimo_numero FROM contadores_codigo WHERE tipo='ALB' AND anio=?", (anio,)).fetchone()
    numero = row["ultimo_numero"]
    return f"ALB-{anio}-{numero:03d}"

def crear_estructura_nextcloud(codigo: str, cliente_nombre: str) -> str:
    nc_url  = os.environ.get("NC_URL", "")
    nc_user = os.environ.get("NC_USER", "")
    nc_pass = os.environ.get("NC_PASS", "")
    nc_base = os.environ.get("NC_BASE_PATH", "/remote.php/dav/files/App/Proyectos")

    if not nc_url:
        # Sin configuración de Nextcloud, devolvemos ruta lógica sin crear nada
        return f"/Proyectos/{codigo} - {cliente_nombre}"

    carpeta = f"{codigo} - {cliente_nombre}"
    carpeta_enc = urllib.parse.quote(carpeta)
    base = f"{nc_url}{nc_base}/{carpeta_enc}"
    auth = (nc_user, nc_pass)
    rutas = [
        base,
        f"{base}/trabajo",
        f"{base}/intercambio",
        f"{base}/intercambio/entrada",
        f"{base}/intercambio/salida",
        f"{base}/fotos",
    ]
    for ruta in rutas:
        try:
            r = requests.request("MKCOL", ruta, auth=auth, verify=False, timeout=10)
            # 201 = creada, 405 = ya existía — ambas son OK
            if r.status_code not in (201, 405):
                print(f"[NC] Warning MKCOL {ruta} → {r.status_code}")
        except Exception as e:
            print(f"[NC] Error creando {ruta}: {e}")

    ruta_base = f"/Proyectos/{carpeta}"
    return ruta_base

@app.post("/proyectos")
async def crear_proyecto(request: Request):
    d = await request.json()
    if d.get("tipo_trabajo") not in TIPOS_TRABAJO:
        raise HTTPException(400, f"tipo_trabajo debe ser uno de {TIPOS_TRABAJO}")
    if d.get("marca") not in MARCAS:
        raise HTTPException(400, f"marca debe ser una de {MARCAS}")

    conn = get_db()
    codigo = generar_codigo(conn, d["tipo_trabajo"])

    cliente = conn.execute("SELECT nombre FROM clientes WHERE id=?", (d.get("cliente_id"),)).fetchone()
    cliente_nombre = cliente["nombre"] if cliente else "SinCliente"
    ruta_nextcloud = crear_estructura_nextcloud(codigo, cliente_nombre)

    cur = conn.execute("""INSERT INTO proyectos
        (codigo, nombre, cliente_id, contacto_id, marca, tipo_trabajo, estado,
         fecha_entrega_compromiso, ruta_nextcloud, notas, creado_por)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (codigo, d.get("nombre"), d.get("cliente_id"), d.get("contacto_id"),
         d["marca"], d["tipo_trabajo"], d.get("estado", "presupuestado"),
         d.get("fecha_entrega_compromiso"), ruta_nextcloud, d.get("notas"), _usuario(request)))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id, "codigo": codigo, "ruta_nextcloud": ruta_nextcloud}

@app.get("/proyectos")
def listar_proyectos(cliente_id: int = None, marca: str = None, tipo_trabajo: str = None, estado: str = None):
    query = """SELECT p.*,
        (SELECT COUNT(*) FROM proyecto_fotos f WHERE f.proyecto_id = p.id) AS num_fotos,
        (SELECT COUNT(*) FROM proyecto_comentarios c WHERE c.proyecto_id = p.id) AS num_comentarios
        FROM proyectos p WHERE 1=1"""
    params = []
    if cliente_id: query += " AND cliente_id=?"; params.append(cliente_id)
    if marca: query += " AND marca=?"; params.append(marca)
    if tipo_trabajo: query += " AND tipo_trabajo=?"; params.append(tipo_trabajo)
    if estado: query += " AND estado=?"; params.append(estado)
    query += " ORDER BY fecha_alta DESC"
    conn = get_db()
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ──────────────────────────────────────────────────────────────────
# FOTOS ADJUNTAS A UN PROYECTO (tarjetas del tablero)
# Se guardan en Nextcloud, dentro de la propia carpeta del proyecto
# (subcarpeta "fotos", igual que "intercambio"), no en el LXC.
# ──────────────────────────────────────────────────────────────────

def _ruta_fotos_nextcloud(proyecto_id: int):
    """URL WebDAV de la subcarpeta 'fotos' dentro de la carpeta del proyecto en Nextcloud.
    None si Nextcloud no está configurado o el proyecto no tiene carpeta asociada."""
    nc_url = os.environ.get("NC_URL", "")
    if not nc_url:
        return None
    conn = get_db()
    row = conn.execute("SELECT ruta_nextcloud FROM proyectos WHERE id=?", (proyecto_id,)).fetchone()
    conn.close()
    if not row or not row["ruta_nextcloud"]:
        return None
    nc_base = os.environ.get("NC_BASE_PATH", "/remote.php/dav/files/App/Proyectos")
    carpeta = row["ruta_nextcloud"].rsplit("/", 1)[-1]  # ruta_nextcloud guardada: "/Proyectos/{codigo} - {cliente}"
    carpeta_enc = urllib.parse.quote(carpeta)
    return f"{nc_url}{nc_base}/{carpeta_enc}/fotos"

def _nc_auth():
    return (os.environ.get("NC_USER", ""), os.environ.get("NC_PASS", ""))

@app.post("/proyectos/{proyecto_id}/fotos")
async def subir_foto_proyecto(proyecto_id: int, request: Request, file: UploadFile = File(...)):
    usuario = _usuario(request)
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in (".jpg", ".jpeg", ".png", ".webp", ".gif"):
        raise HTTPException(400, "Formato de imagen no soportado (usa jpg, png, webp o gif)")
    carpeta_fotos = _ruta_fotos_nextcloud(proyecto_id)
    if not carpeta_fotos:
        raise HTTPException(400, "Este proyecto no tiene carpeta de Nextcloud asociada (o Nextcloud no está configurado)")
    nombre_archivo = f"{secrets.token_hex(8)}{ext}"  # el nombre original no se conserva, no hace falta
    contenido = await file.read()
    auth = _nc_auth()
    try:
        # por si el proyecto se creó antes de que existiera esta subcarpeta
        requests.request("MKCOL", carpeta_fotos, auth=auth, verify=False, timeout=10)
        r = requests.request("PUT", f"{carpeta_fotos}/{urllib.parse.quote(nombre_archivo)}",
                              data=contenido, auth=auth, verify=False, timeout=30)
        if r.status_code not in (200, 201, 204):
            raise HTTPException(502, f"Nextcloud rechazó la subida (HTTP {r.status_code})")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(502, f"Error subiendo la foto a Nextcloud: {e}")
    conn = get_db()
    cur = conn.execute("INSERT INTO proyecto_fotos (proyecto_id, filename, usuario) VALUES (?,?,?)",
                        (proyecto_id, nombre_archivo, usuario))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id, "filename": nombre_archivo}

@app.get("/proyectos/{proyecto_id}/fotos")
def listar_fotos_proyecto(proyecto_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM proyecto_fotos WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/proyectos/{proyecto_id}/fotos/{filename}")
def obtener_foto_proyecto(proyecto_id: int, filename: str):
    filename = os.path.basename(filename)  # el nombre siempre lo generamos nosotros, pero por si acaso
    carpeta_fotos = _ruta_fotos_nextcloud(proyecto_id)
    if not carpeta_fotos:
        raise HTTPException(404, "Foto no encontrada")
    try:
        r = requests.get(f"{carpeta_fotos}/{urllib.parse.quote(filename)}", auth=_nc_auth(), verify=False, timeout=20)
    except Exception as e:
        raise HTTPException(502, f"No se pudo contactar con Nextcloud: {e}")
    if r.status_code != 200:
        raise HTTPException(404, "Foto no encontrada en Nextcloud")
    return StreamingResponse(io.BytesIO(r.content), media_type=r.headers.get("Content-Type", "image/jpeg"))

@app.delete("/proyectos/{proyecto_id}/fotos/{foto_id}")
def borrar_foto_proyecto(proyecto_id: int, foto_id: int):
    conn = get_db()
    row = conn.execute("SELECT * FROM proyecto_fotos WHERE id=? AND proyecto_id=?", (foto_id, proyecto_id)).fetchone()
    if row:
        carpeta_fotos = _ruta_fotos_nextcloud(proyecto_id)
        if carpeta_fotos:
            try:
                requests.request("DELETE", f"{carpeta_fotos}/{urllib.parse.quote(row['filename'])}",
                                  auth=_nc_auth(), verify=False, timeout=10)
            except Exception:
                pass
        conn.execute("DELETE FROM proyecto_fotos WHERE id=?", (foto_id,))
        conn.commit()
    conn.close()
    return {"status": "deleted"}

# ──────────────────────────────────────────────────────────────────
# COMENTARIOS EN UN PROYECTO (tarjetas del tablero)
# ──────────────────────────────────────────────────────────────────

@app.post("/proyectos/{proyecto_id}/comentarios")
async def crear_comentario(proyecto_id: int, request: Request):
    usuario = _usuario(request)
    d = await request.json()
    texto = (d.get("texto") or "").strip()
    if not texto:
        raise HTTPException(400, "El comentario está vacío")
    conn = get_db()
    cur = conn.execute("INSERT INTO proyecto_comentarios (proyecto_id, usuario, texto) VALUES (?,?,?)",
                        (proyecto_id, usuario, texto))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.get("/proyectos/{proyecto_id}/comentarios")
def listar_comentarios_proyecto(proyecto_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM proyecto_comentarios WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.delete("/proyectos/{proyecto_id}/comentarios/{comentario_id}")
def borrar_comentario_proyecto(proyecto_id: int, comentario_id: int):
    conn = get_db()
    conn.execute("DELETE FROM proyecto_comentarios WHERE id=? AND proyecto_id=?", (comentario_id, proyecto_id))
    conn.commit()
    conn.close()
    return {"status": "deleted"}

@app.get("/proyectos/{proyecto_id}")
def detalle_proyecto(proyecto_id: int):
    conn = get_db()
    proyecto = conn.execute("SELECT * FROM proyectos WHERE id=?", (proyecto_id,)).fetchone()
    if not proyecto:
        conn.close()
        raise HTTPException(404, "Proyecto no encontrado")
    presupuestos = conn.execute("SELECT * FROM presupuestos WHERE proyecto_id=? ORDER BY version", (proyecto_id,)).fetchall()
    facturas = conn.execute("SELECT * FROM facturas WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    horas = conn.execute("SELECT * FROM horas WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    gastos = conn.execute("SELECT * FROM gastos WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    conn.close()
    return {
        "proyecto": dict(proyecto),
        "presupuestos": [dict(r) for r in presupuestos],
        "facturas": [dict(r) for r in facturas],
        "horas": [dict(r) for r in horas],
        "gastos": [dict(r) for r in gastos],
    }

@app.patch("/proyectos/{proyecto_id}")
async def editar_proyecto(proyecto_id: int, request: Request):
    d = await request.json()
    campos_permitidos = ["nombre", "estado", "fecha_entrega_compromiso", "fecha_entrega_real",
                          "fecha_facturacion", "fecha_cobro", "contacto_id", "prioridad"]
    sets, params = [], []
    for campo in campos_permitidos:
        if campo in d:
            sets.append(f"{campo}=?")
            params.append(d[campo])
    if not sets:
        raise HTTPException(400, "Nada que actualizar")
    params.append(proyecto_id)
    conn = get_db()
    conn.execute(f"UPDATE proyectos SET {', '.join(sets)} WHERE id=?", params)
    conn.commit()
    conn.close()
    return {"status": "updated"}

@app.patch("/proyectos/{proyecto_id}/imagen")
async def actualizar_imagen(proyecto_id: int, request: Request):
    d = await request.json()
    # STUB: la subida real del archivo a Nextcloud (trabajo/) se conecta más adelante
    conn = get_db()
    conn.execute("UPDATE proyectos SET imagen_portada_ruta=? WHERE id=?", (d.get("ruta"), proyecto_id))
    conn.commit()
    conn.close()
    return {"status": "updated"}

@app.patch("/proyectos/{proyecto_id}/notas")
async def actualizar_notas(proyecto_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    conn.execute("UPDATE proyectos SET notas=? WHERE id=?", (d.get("notas"), proyecto_id))
    conn.commit()
    conn.close()
    return {"status": "updated"}

# ──────────────────────────────────────────────────────────────────
# PRESUPUESTOS / FACTURAS / HORAS (por proyecto)
# ──────────────────────────────────────────────────────────────────

@app.get("/tarifas")
def listar_tarifas():
    conn = get_db()
    rows = conn.execute("SELECT * FROM tarifas_horas").fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/tarifas")
async def actualizar_tarifa(request: Request):
    d = await request.json()
    conn = get_db()
    conn.execute("""INSERT INTO tarifas_horas (tipo_hora, precio_hora) VALUES (?,?)
        ON CONFLICT(tipo_hora) DO UPDATE SET precio_hora=excluded.precio_hora""",
        (d["tipo_hora"], d["precio_hora"]))
    conn.commit()
    conn.close()
    return {"status": "saved"}

@app.post("/proyectos/{proyecto_id}/presupuestos")
async def crear_presupuesto(proyecto_id: int, request: Request):
    d = await request.json() if await request.body() else {}
    conn = get_db()
    ultima = conn.execute("SELECT MAX(version) as v FROM presupuestos WHERE proyecto_id=?", (proyecto_id,)).fetchone()
    version = (ultima["v"] or 0) + 1
    cur = conn.execute("INSERT INTO presupuestos (proyecto_id, version, estado, creado_por) VALUES (?,?,?,?)",
                        (proyecto_id, version, d.get("estado", "borrador"), _usuario(request)))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id, "version": version}

@app.get("/proyectos/{proyecto_id}/presupuestos")
def listar_presupuestos(proyecto_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM presupuestos WHERE proyecto_id=? ORDER BY version", (proyecto_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/presupuestos/{presupuesto_id}")
def detalle_presupuesto(presupuesto_id: int):
    conn = get_db()
    presupuesto = conn.execute("SELECT * FROM presupuestos WHERE id=?", (presupuesto_id,)).fetchone()
    if not presupuesto:
        conn.close()
        raise HTTPException(404, "Presupuesto no encontrado")
    lineas = conn.execute("SELECT * FROM presupuesto_lineas WHERE presupuesto_id=? ORDER BY id", (presupuesto_id,)).fetchall()
    conn.close()
    return {"presupuesto": dict(presupuesto), "lineas": [dict(r) for r in lineas]}

def _recalcular_total_presupuesto(conn, presupuesto_id: int):
    total = conn.execute("SELECT COALESCE(SUM(importe),0) as t FROM presupuesto_lineas WHERE presupuesto_id=?",
                          (presupuesto_id,)).fetchone()["t"]
    conn.execute("UPDATE presupuestos SET importe=? WHERE id=?", (total, presupuesto_id))
    return total

@app.post("/presupuestos/{presupuesto_id}/lineas")
async def crear_linea_presupuesto(presupuesto_id: int, request: Request):
    d = await request.json()
    cantidad = float(d.get("cantidad", 1))
    precio_unitario = float(d.get("precio_unitario", 0))
    importe = round(cantidad * precio_unitario, 2)
    conn = get_db()
    cur = conn.execute("""INSERT INTO presupuesto_lineas (presupuesto_id, tipo, descripcion, cantidad, precio_unitario, importe)
        VALUES (?,?,?,?,?,?)""",
        (presupuesto_id, d.get("tipo"), d.get("descripcion"), cantidad, precio_unitario, importe))
    conn.commit()
    new_id = cur.lastrowid
    total = _recalcular_total_presupuesto(conn, presupuesto_id)
    conn.commit()
    conn.close()
    return {"status": "saved", "id": new_id, "importe": importe, "total_presupuesto": total}

@app.delete("/presupuestos/{presupuesto_id}/lineas/{linea_id}")
def borrar_linea_presupuesto(presupuesto_id: int, linea_id: int):
    conn = get_db()
    conn.execute("DELETE FROM presupuesto_lineas WHERE id=? AND presupuesto_id=?", (linea_id, presupuesto_id))
    conn.commit()
    total = _recalcular_total_presupuesto(conn, presupuesto_id)
    conn.commit()
    conn.close()
    return {"status": "deleted", "total_presupuesto": total}

@app.patch("/presupuestos/{presupuesto_id}")
async def editar_presupuesto(presupuesto_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    if "estado" in d:
        conn.execute("UPDATE presupuestos SET estado=? WHERE id=?", (d["estado"], presupuesto_id))
    conn.commit()
    conn.close()
    return {"status": "updated"}

# ──────────────────────────────────────────────────────────────────
# ALBARANES DE ENTREGA
# ──────────────────────────────────────────────────────────────────

@app.post("/proyectos/{proyecto_id}/albaranes")
async def crear_albaran(proyecto_id: int, request: Request):
    usuario = _usuario(request)
    d = await request.json() if await request.body() else {}
    presupuesto_id = d.get("presupuesto_id")
    lineas_manuales = d.get("lineas")  # opcional: [{descripcion, cantidad}, ...]

    conn = get_db()
    proyecto = conn.execute("SELECT id FROM proyectos WHERE id=?", (proyecto_id,)).fetchone()
    if not proyecto:
        conn.close()
        raise HTTPException(404, "Proyecto no encontrado")

    numero = generar_numero_albaran(conn)
    cur = conn.execute("""INSERT INTO albaranes (proyecto_id, presupuesto_id, numero, estado, notas, creado_por)
        VALUES (?,?,?,?,?,?)""",
        (proyecto_id, presupuesto_id, numero, "emitido", d.get("notas"), usuario))
    albaran_id = cur.lastrowid

    if lineas_manuales:
        for linea in lineas_manuales:
            conn.execute("INSERT INTO albaran_lineas (albaran_id, descripcion, cantidad) VALUES (?,?,?)",
                         (albaran_id, linea.get("descripcion"), linea.get("cantidad")))
    elif presupuesto_id:
        # heredar las líneas del presupuesto aceptado, tal cual pediste
        lineas_presu = conn.execute("SELECT descripcion, cantidad FROM presupuesto_lineas WHERE presupuesto_id=? ORDER BY id",
                                     (presupuesto_id,)).fetchall()
        for l in lineas_presu:
            conn.execute("INSERT INTO albaran_lineas (albaran_id, descripcion, cantidad) VALUES (?,?,?)",
                         (albaran_id, l["descripcion"], l["cantidad"]))

    conn.commit()
    conn.close()
    return {"status": "saved", "id": albaran_id, "numero": numero}

@app.get("/proyectos/{proyecto_id}/albaranes")
def listar_albaranes_proyecto(proyecto_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM albaranes WHERE proyecto_id=? ORDER BY fecha DESC", (proyecto_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/albaranes/{albaran_id}")
def detalle_albaran(albaran_id: int):
    conn = get_db()
    albaran = conn.execute("""
        SELECT a.*, p.codigo AS proyecto_codigo, p.nombre AS proyecto_nombre, p.marca,
               c.nombre AS cliente_nombre, ct.nombre AS contacto_nombre
        FROM albaranes a
        JOIN proyectos p ON p.id = a.proyecto_id
        LEFT JOIN clientes c ON c.id = p.cliente_id
        LEFT JOIN contactos ct ON ct.id = p.contacto_id
        WHERE a.id = ?
    """, (albaran_id,)).fetchone()
    if not albaran:
        conn.close()
        raise HTTPException(404, "Albarán no encontrado")
    lineas = conn.execute("SELECT * FROM albaran_lineas WHERE albaran_id=? ORDER BY id", (albaran_id,)).fetchall()
    conn.close()
    return {"albaran": dict(albaran), "lineas": [dict(l) for l in lineas]}

@app.patch("/albaranes/{albaran_id}")
async def editar_albaran(albaran_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    if "estado" in d:
        conn.execute("UPDATE albaranes SET estado=? WHERE id=?", (d["estado"], albaran_id))
    if "notas" in d:
        conn.execute("UPDATE albaranes SET notas=? WHERE id=?", (d["notas"], albaran_id))
    conn.commit()
    conn.close()
    return {"status": "updated"}

@app.delete("/albaranes/{albaran_id}")
def borrar_albaran(albaran_id: int):
    conn = get_db()
    conn.execute("DELETE FROM albaran_lineas WHERE albaran_id=?", (albaran_id,))
    conn.execute("DELETE FROM albaranes WHERE id=?", (albaran_id,))
    conn.commit()
    conn.close()
    return {"status": "deleted"}

@app.post("/albaranes/{albaran_id}/lineas")
async def crear_linea_albaran(albaran_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    cur = conn.execute("INSERT INTO albaran_lineas (albaran_id, descripcion, cantidad) VALUES (?,?,?)",
                        (albaran_id, d.get("descripcion"), d.get("cantidad")))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.delete("/albaranes/{albaran_id}/lineas/{linea_id}")
def borrar_linea_albaran(albaran_id: int, linea_id: int):
    conn = get_db()
    conn.execute("DELETE FROM albaran_lineas WHERE id=? AND albaran_id=?", (linea_id, albaran_id))
    conn.commit()
    conn.close()
    return {"status": "deleted"}

@app.post("/proyectos/{proyecto_id}/facturas")
async def crear_factura(proyecto_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    cur = conn.execute("""INSERT INTO facturas (proyecto_id, numero, importe, fecha, estado_cobro, referencia_archivo)
        VALUES (?,?,?,?,?,?)""",
        (proyecto_id, d.get("numero"), d.get("importe"), d.get("fecha"), d.get("estado_cobro"), d.get("referencia_archivo")))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.get("/proyectos/{proyecto_id}/facturas")
def listar_facturas(proyecto_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM facturas WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/proyectos/{proyecto_id}/horas")
async def crear_hora(proyecto_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    cur = conn.execute("""INSERT INTO horas (proyecto_id, fecha, tipo_trabajo_hora, horas, descripcion, usuario)
        VALUES (?,?,?,?,?,?)""",
        (proyecto_id, d.get("fecha"), d.get("tipo_trabajo_hora"), d.get("horas"), d.get("descripcion"), _usuario(request)))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.get("/proyectos/{proyecto_id}/horas")
def listar_horas(proyecto_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM horas WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/proyectos/{proyecto_id}/gastos")
async def crear_gasto(proyecto_id: int, request: Request):
    d = await request.json()
    conn = get_db()
    cur = conn.execute("""INSERT INTO gastos (proyecto_id, fecha, descripcion, importe, usuario)
        VALUES (?,?,?,?,?)""",
        (proyecto_id, d.get("fecha"), d.get("descripcion"), d.get("importe"), _usuario(request)))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return {"status": "saved", "id": new_id}

@app.get("/proyectos/{proyecto_id}/gastos")
def listar_gastos(proyecto_id: int):
    conn = get_db()
    rows = conn.execute("SELECT * FROM gastos WHERE proyecto_id=? ORDER BY fecha", (proyecto_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ──────────────────────────────────────────────────────────────────
# TABLERO KANBAN
# ──────────────────────────────────────────────────────────────────

@app.get("/tablero")
def tablero():
    conn = get_db()
    rows = conn.execute("""
        SELECT p.*, c.nombre as cliente_nombre
        FROM proyectos p LEFT JOIN clientes c ON p.cliente_id = c.id
        WHERE p.estado != 'rechazado'
        ORDER BY p.fecha_alta DESC
    """).fetchall()
    conn.close()
    columnas = {estado: [] for estado in ESTADOS_PROYECTO if estado != "rechazado"}
    for r in rows:
        columnas.setdefault(r["estado"], []).append(dict(r))
    return columnas

# ──────────────────────────────────────────────────────────
# EXPORTAR PRESUPUESTO A WORD
# ──────────────────────────────────────────────────────────

from fastapi.responses import StreamingResponse
from docx import Document as DocxDocument
from docx.shared import RGBColor

TIPO_HORA_LABEL = {
    'ir':             'Ingeniería inversa',
    'diseno':         'Diseño CAD',
    'diseno_organico':'Diseño orgánico',
    'impresion':      'Impresión 3D',
    'reunion':        'Reunión',
    'montaje':        'Montaje',
    'km':             'Kilometraje',
    'impresion_3d':   'Impresión 3D',
}

def _label_tipo(tipo_raw: str) -> str:
    """Convierte 'horas_ir' o 'impresion_3d' en texto legible."""
    t = tipo_raw.replace('horas_', '').strip()
    return TIPO_HORA_LABEL.get(t, t.replace('_', ' ').capitalize())

def _reemplazar_en_parrafo(paragraph, reemplazos: dict):
    """Sustituye texto en un párrafo (puede estar repartido en varios runs),
    conservando el formato del primer run."""
    texto = ''.join(r.text for r in paragraph.runs)
    if not texto:
        return
    nuevo = texto
    cambiado = False
    for clave, valor in reemplazos.items():
        if clave in nuevo:
            nuevo = nuevo.replace(clave, valor)
            cambiado = True
    if cambiado and paragraph.runs:
        paragraph.runs[0].text = nuevo
        for r in paragraph.runs[1:]:
            r.text = ''


def _set_cell_text(cell, text: str):
    """Sustituye el texto de una celda conservando el formato del primer run existente."""
    paragraphs = cell.paragraphs
    para = paragraphs[0] if paragraphs else cell.add_paragraph()
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra.text = ''
    else:
        para.add_run(text)
    for p in paragraphs[1:]:
        p._element.getparent().remove(p._element)


def _clonar_fila_antes(tabla, fila_modelo_xml, fila_referencia):
    """Clona el XML de una fila de plantilla (conserva sus anchos de columna y formato
    exactos) y la inserta justo antes de fila_referencia. add_row() no sirve para esto:
    reparte el ancho de las columnas de forma uniforme y descuadra la tabla."""
    import copy
    from docx.table import _Row
    nueva_tr = copy.deepcopy(fila_modelo_xml)
    fila_referencia._tr.addprevious(nueva_tr)
    return _Row(nueva_tr, tabla)


# ──────────────────────────────────────────────────────────
# PLANTILLAS DE PRESUPUESTO POR MARCA (Nextcloud, con caché local)
# ──────────────────────────────────────────────────────────

NC_PLANTILLAS_PATH = os.environ.get("NC_PLANTILLAS_PATH", "/remote.php/dav/files/App/Plantillas/app")
PLANTILLAS_CACHE_DIR = "/opt/print3d/plantillas_cache"

PLANTILLA_POR_MARCA = {
    "myrox_lab":        "Plantilla_Myrox_Lab.docx",
    "myrox_print":      "Plantilla_Myrox_Print.docx",
    "myrox_works":      "Plantilla_Myrox_Works.docx",
    "myrox_automation": "Plantilla_Myrox_Automation.docx",
    "mw3d":             "Plantilla_MW3D_Studio.docx",
}


def _descargar_plantilla_presupuesto(marca: str) -> bytes:
    """Descarga de Nextcloud la plantilla Word que corresponde a la marca del proyecto.
    Si Nextcloud falla, cae a una copia en caché local guardada en la última descarga OK."""
    filename = PLANTILLA_POR_MARCA.get(marca)
    if not filename:
        raise ValueError(f"No hay plantilla configurada para la marca '{marca}'")

    nc_url  = os.environ.get("NC_URL", "")
    nc_user = os.environ.get("NC_USER", "")
    nc_pass = os.environ.get("NC_PASS", "")
    cache_path = os.path.join(PLANTILLAS_CACHE_DIR, filename)

    if nc_url:
        try:
            ruta = f"{nc_url}{NC_PLANTILLAS_PATH}/{urllib.parse.quote(filename)}"
            r = requests.get(ruta, auth=(nc_user, nc_pass), verify=False, timeout=15)
            if r.status_code == 200:
                try:
                    os.makedirs(PLANTILLAS_CACHE_DIR, exist_ok=True)
                    with open(cache_path, "wb") as f:
                        f.write(r.content)
                except Exception as e:
                    print(f"[NC] No se pudo actualizar la caché de {filename}: {e}")
                return r.content
            print(f"[NC] Plantilla {filename} → HTTP {r.status_code}, probando caché local")
        except Exception as e:
            print(f"[NC] Error descargando plantilla {filename} de Nextcloud: {e}")

    if os.path.exists(cache_path):
        with open(cache_path, "rb") as f:
            return f.read()

    raise FileNotFoundError(
        f"No se pudo obtener la plantilla '{filename}' (falló Nextcloud y no hay caché local en {cache_path})"
    )


def _build_presupuesto_doc(data: dict, marca: str, tipo_trabajo: str) -> bytes:
    """Rellena la plantilla Word de la marca del proyecto con los datos del presupuesto
    y devuelve los bytes del .docx resultante. El diseño (logo, colores, título, columnas
    de la tabla de trabajo) viene ya hecho en cada plantilla — aquí solo se insertan datos."""
    from datetime import datetime as _dt

    plantilla_bytes = _descargar_plantilla_presupuesto(marca)
    doc = DocxDocument(io.BytesIO(plantilla_bytes))

    proyecto = data['proyecto']
    lineas   = data['lineas']
    cliente_nombre  = proyecto.get('cliente_nombre') or ''
    contacto_nombre = proyecto.get('contacto_nombre') or ''
    codigo = proyecto.get('codigo') or ''

    fecha_raw = proyecto.get('fecha_alta', '')[:10] if proyecto.get('fecha_alta') else ''
    try:
        fecha = _dt.strptime(fecha_raw, '%Y-%m-%d').strftime('%d/%m/%Y')
    except Exception:
        fecha = ''

    reemplazos = {'[CÓDIGO]': codigo, '[FECHA]': fecha}

    # --- Cabecera: sustituir [CÓDIGO] y [FECHA] estén donde estén ---
    for p in doc.paragraphs:
        _reemplazar_en_parrafo(p, reemplazos)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _reemplazar_en_parrafo(p, reemplazos)

    # --- Tabla "Datos del proyecto" (Nombre / Cliente / Contacto / Fecha) ---
    tabla_datos = None
    for t in doc.tables:
        etiquetas = [r.cells[0].text.strip() for r in t.rows[:4]]
        if etiquetas == ['Nombre del proyecto', 'Cliente / Empresa', 'Persona de contacto', 'Fecha']:
            tabla_datos = t
            break
    if tabla_datos:
        _set_cell_text(tabla_datos.rows[0].cells[1], proyecto.get('nombre') or '')
        _set_cell_text(tabla_datos.rows[1].cells[1], cliente_nombre)
        _set_cell_text(tabla_datos.rows[2].cells[1], contacto_nombre)
        _set_cell_text(tabla_datos.rows[3].cells[1], fecha)

    # --- Descripción general del encargo (si el proyecto tiene notas) ---
    notas = (proyecto.get('notas') or '').strip()
    if notas:
        for p in doc.paragraphs:
            if p.text.strip().startswith('[Describir') or p.text.strip().startswith('['):
                # Sustituye el párrafo entero por las notas del proyecto
                if p.runs:
                    p.runs[0].text = notas
                    for r in p.runs[1:]:
                        r.text = ''
                break

    # --- Tabla de líneas de trabajo: se identifica porque su última columna es "Importe" ---
    tabla_lineas = None
    for t in doc.tables:
        cabecera = [c.text.strip() for c in t.rows[0].cells]
        if cabecera and cabecera[-1] == 'Importe' and len(t.rows) >= 4:
            tabla_lineas = t
            break

    subtotal = round(sum(l.get('importe', 0) or 0 for l in lineas), 2)
    iva      = round(subtotal * 0.21, 2)
    total    = round(subtotal + iva, 2)

    if tabla_lineas:
        import copy
        filas = tabla_lineas.rows
        subtotal_row, iva_row, total_row = filas[-3], filas[-2], filas[-1]
        content_rows = filas[1:len(filas) - 3]  # filas vacías de plantilla para líneas
        # Copia "limpia" (aún sin rellenar) de una fila de plantilla, para clonarla si hacen falta más
        fila_modelo_xml = copy.deepcopy(content_rows[-1]._tr) if content_rows else None

        for i, linea in enumerate(lineas):
            if i < len(content_rows):
                row = content_rows[i]
            elif fila_modelo_xml is not None:
                row = _clonar_fila_antes(tabla_lineas, fila_modelo_xml, subtotal_row)
            else:
                row = tabla_lineas.add_row()  # última red de seguridad si la plantilla no tenía filas vacías
            desc = linea.get('descripcion') or ''
            cantidad = linea.get('cantidad')
            if cantidad and cantidad != 1:
                desc = f"{desc} (x{cantidad:g})"
            _set_cell_text(row.cells[0], desc)
            _set_cell_text(row.cells[-1], f"{(linea.get('importe') or 0):.2f} €")

        # borrar filas de plantilla que sobren si hay menos líneas que huecos
        for row in content_rows[len(lineas):]:
            row._tr.getparent().remove(row._tr)

        _set_cell_text(subtotal_row.cells[-1], f"{subtotal:.2f} €")
        _set_cell_text(iva_row.cells[-1], f"{iva:.2f} €")
        _set_cell_text(total_row.cells[-1], f"{total:.2f} €")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@app.get("/presupuestos/{presupuesto_id}/exportar")
def exportar_presupuesto(presupuesto_id: int):
    conn = get_db()
    presupuesto = conn.execute("SELECT * FROM presupuestos WHERE id=?", (presupuesto_id,)).fetchone()
    if not presupuesto:
        conn.close()
        raise HTTPException(404, "Presupuesto no encontrado")

    lineas = conn.execute("SELECT * FROM presupuesto_lineas WHERE presupuesto_id=? ORDER BY id", (presupuesto_id,)).fetchall()
    proyecto = conn.execute("""
        SELECT p.*, c.nombre as cliente_nombre, ct.nombre as contacto_nombre
        FROM proyectos p
        LEFT JOIN clientes c ON p.cliente_id = c.id
        LEFT JOIN contactos ct ON p.contacto_id = ct.id
        WHERE p.id = ?
    """, (presupuesto['proyecto_id'],)).fetchone()
    conn.close()

    if not proyecto:
        raise HTTPException(404, "Proyecto no encontrado")

    tipo_trabajo = proyecto['tipo_trabajo']
    marca        = proyecto['marca'] or 'myrox_lab'
    if marca not in PLANTILLA_POR_MARCA:
        raise HTTPException(400, f"Marca de proyecto desconocida: '{marca}'")

    data = {
        'proyecto': dict(proyecto),
        'lineas':   [dict(l) for l in lineas],
    }

    try:
        docx_bytes = _build_presupuesto_doc(data, marca, tipo_trabajo)
    except FileNotFoundError as e:
        raise HTTPException(500, f"No se pudo cargar la plantilla de '{marca}': {e}")
    except Exception as e:
        raise HTTPException(500, f"Error generando el documento: {e}")

    codigo = proyecto['codigo'] or f"presupuesto_{presupuesto_id}"
    filename = f"Presupuesto_{codigo}_v{presupuesto['version']}.docx"

    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# ──────────────────────────────────────────────────────────────────
# EXPORTAR ALBARÁN A WORD (construido en código, con el logo de la marca)
# ──────────────────────────────────────────────────────────────────

LOGOS_DIR = "/opt/print3d/logos"
LOGO_POR_MARCA = {
    "myrox_lab":        "logo_myrox_lab.png",
    "myrox_print":      "logo_myrox_print.png",
    "myrox_works":      "logo_myrox_works.png",
    "myrox_automation": "logo_myrox_automation.png",
    "mw3d":             "logo_mw3d.png",
}
ALBARAN_ACCENT = RGBColor(0x1A, 0x3A, 0x5C)
ALBARAN_GREY = RGBColor(0x6B, 0x72, 0x80)

def _logo_bytes_marca(marca: str) -> bytes:
    filename = LOGO_POR_MARCA.get(marca)
    if not filename:
        raise ValueError(f"No hay logo configurado para la marca '{marca}'")
    ruta = os.path.join(LOGOS_DIR, filename)
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"Falta el logo de '{marca}' en {ruta}")
    with open(ruta, "rb") as f:
        return f.read()

def _build_albaran_doc(data: dict) -> bytes:
    """Construye el Word del albarán en código (sin plantilla), con el logo de la marca del proyecto."""
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from datetime import datetime as _dt

    albaran = data["albaran"]
    lineas = data["lineas"]
    marca = albaran.get("marca") or "myrox_lab"

    doc = DocxDocument()
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)

    # --- Cabecera: logo + título + nº/fecha ---
    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = True
    logo_cell, info_cell = header_table.rows[0].cells
    try:
        logo_bytes = _logo_bytes_marca(marca)
        logo_cell.paragraphs[0].add_run().add_picture(io.BytesIO(logo_bytes), width=Cm(4))
    except (FileNotFoundError, ValueError):
        logo_cell.paragraphs[0].add_run(marca.upper())

    p_titulo = info_cell.paragraphs[0]
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_titulo.add_run("ALBARÁN DE ENTREGA")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = ALBARAN_ACCENT

    fecha_raw = (albaran.get("fecha") or "")[:10]
    try:
        fecha = _dt.strptime(fecha_raw, "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        fecha = fecha_raw

    p_meta = info_cell.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rm = p_meta.add_run(f"Nº {albaran['numero']}   ·   Fecha: {fecha}")
    rm.font.size = Pt(10); rm.font.color.rgb = ALBARAN_GREY

    doc.add_paragraph()

    # --- Datos de cliente / proyecto ---
    datos_tabla = doc.add_table(rows=0, cols=2)
    for etiqueta, valor in [
        ("Cliente", albaran.get("cliente_nombre") or "—"),
        ("Contacto", albaran.get("contacto_nombre") or "—"),
        ("Proyecto", f"{albaran.get('proyecto_codigo') or ''} — {albaran.get('proyecto_nombre') or ''}"),
    ]:
        row = datos_tabla.add_row()
        row.cells[0].paragraphs[0].add_run(etiqueta).bold = True
        row.cells[1].paragraphs[0].add_run(valor)

    doc.add_paragraph()

    # --- Tabla de líneas entregadas ---
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = "Light Grid Accent 1"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tabla.rows[0].cells
    hdr[0].paragraphs[0].add_run("Descripción").bold = True
    hdr[1].paragraphs[0].add_run("Cantidad").bold = True
    for linea in lineas:
        row = tabla.add_row()
        row.cells[0].paragraphs[0].add_run(linea.get("descripcion") or "")
        cant = linea.get("cantidad")
        row.cells[1].paragraphs[0].add_run(f"{cant:g}" if cant is not None else "")

    doc.add_paragraph()

    if albaran.get("notas"):
        p_notas = doc.add_paragraph()
        p_notas.add_run("Observaciones: ").bold = True
        p_notas.add_run(albaran["notas"])
        doc.add_paragraph()

    # --- Firma ---
    firma_tabla = doc.add_table(rows=2, cols=2)
    firma_tabla.rows[0].cells[0].paragraphs[0].add_run("Entregado por:").bold = True
    firma_tabla.rows[0].cells[1].paragraphs[0].add_run("Recibido por (nombre y firma):").bold = True
    for cell in firma_tabla.rows[1].cells:
        cell.paragraphs[0].add_run("\n\n\n_______________________")

    doc.add_paragraph()
    p_legal = doc.add_paragraph()
    r_legal = p_legal.add_run(
        "Este albarán acredita la entrega/recepción del material o trabajo descrito y no constituye factura."
    )
    r_legal.italic = True; r_legal.font.size = Pt(8); r_legal.font.color.rgb = ALBARAN_GREY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

@app.get("/albaranes/{albaran_id}/exportar")
def exportar_albaran(albaran_id: int):
    conn = get_db()
    albaran = conn.execute("""
        SELECT a.*, p.codigo AS proyecto_codigo, p.nombre AS proyecto_nombre, p.marca,
               c.nombre AS cliente_nombre, ct.nombre AS contacto_nombre
        FROM albaranes a
        JOIN proyectos p ON p.id = a.proyecto_id
        LEFT JOIN clientes c ON c.id = p.cliente_id
        LEFT JOIN contactos ct ON ct.id = p.contacto_id
        WHERE a.id = ?
    """, (albaran_id,)).fetchone()
    if not albaran:
        conn.close()
        raise HTTPException(404, "Albarán no encontrado")
    lineas = conn.execute("SELECT * FROM albaran_lineas WHERE albaran_id=? ORDER BY id", (albaran_id,)).fetchall()
    conn.close()

    try:
        docx_bytes = _build_albaran_doc({"albaran": dict(albaran), "lineas": [dict(l) for l in lineas]})
    except Exception as e:
        raise HTTPException(500, f"Error generando el albarán: {e}")

    filename = f"Albaran_{albaran['numero']}.docx"
    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# ──────────────────────────────────────────────────────────
# ARCHIVAR / DESARCHIVAR / BORRAR PROYECTO
# ──────────────────────────────────────────────────────────

@app.patch("/proyectos/{proyecto_id}/archivar")
def archivar_proyecto(proyecto_id: int):
    conn = get_db()
    conn.execute("UPDATE proyectos SET archivado=1 WHERE id=?", (proyecto_id,))
    conn.commit(); conn.close()
    return {"status": "archivado"}

@app.patch("/proyectos/{proyecto_id}/desarchivar")
def desarchivar_proyecto(proyecto_id: int):
    conn = get_db()
    conn.execute("UPDATE proyectos SET archivado=0 WHERE id=?", (proyecto_id,))
    conn.commit(); conn.close()
    return {"status": "desarchivado"}

@app.get("/proyectos/archivados")
def listar_archivados():
    conn = get_db()
    rows = conn.execute("""
        SELECT p.*, c.nombre as cliente_nombre
        FROM proyectos p LEFT JOIN clientes c ON p.cliente_id = c.id
        WHERE p.archivado=1 ORDER BY p.fecha_alta DESC
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.delete("/proyectos/{proyecto_id}")
def borrar_proyecto(proyecto_id: int):
    conn = get_db()
    # Solo permite borrar proyectos archivados
    p = conn.execute("SELECT archivado FROM proyectos WHERE id=?", (proyecto_id,)).fetchone()
    if not p:
        conn.close(); raise HTTPException(404, "Proyecto no encontrado")
    if not p["archivado"]:
        conn.close(); raise HTTPException(400, "Solo se pueden borrar proyectos archivados")
    conn.execute("DELETE FROM proyectos WHERE id=?", (proyecto_id,))
    conn.commit(); conn.close()
    return {"status": "eliminado"}
